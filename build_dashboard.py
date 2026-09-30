"""
Rebuilds the "Drop-off Recovery Desk" dashboard HTML from a fresh export of
Wiom_Dropoff_Calling_Pack_03Sep2026.xlsx.

Usage:
    py build_dashboard.py <path-to-downloaded-xlsx> [output-html-path]

All numbers are computed directly from the 'CALL SHEET - APP Drop offs' and
'DAILY SUMMARY -APP Drop offs ' sheets, using the same logic (case-insensitive
COUNTIF/COUNTIFS-style matching, MEDIAN/AVERAGE/PERCENTILE) as the
'Executive Summary - APP Dropoff' tab. See README.md in this folder for the
validation notes.
"""
import sys
import warnings
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from jinja2 import Template

from voc_analysis import classify_voc

warnings.filterwarnings("ignore")

HERE = Path(__file__).parent
TEMPLATE_PATH = HERE / "template.html.j2"

CALL_SHEET = "CALL SHEET - APP Drop offs"
DAILY_SUMMARY = "DAILY SUMMARY -APP Drop offs "

BREAKDOWN_BUCKETS = [
    ("Wouldn't", "cat-1"),
    ("Not Yet", "cat-2"),
    ("Couldn't", "cat-3"),
    ("Exisiting Wiom/Wiom Net Customer", "cat-6"),
]
BREAKDOWN_DISPLAY_NAME = {
    "Exisiting Wiom/Wiom Net Customer": "Existing Wiom / Wiom Net Customer",
}

INTENT_ORDER = [
    ("Wouldn't", "cat-1"),
    ("Not Yet", "cat-2"),
    ("Couldn't", "cat-3"),
    ("Unknown", "cat-4"),
    ("Already Booked", "cat-5"),
    ("Exisiting Wiom/Wiom Net Customer", "cat-6"),
]
INTENT_DISPLAY_NAME = {
    "Exisiting Wiom/Wiom Net Customer": "Existing Wiom/Net",
}


def col(df, letter):
    idx = openpyxl.utils.column_index_from_string(letter.upper()) - 1
    return df.iloc[:, idx]


def norm(s):
    return s.astype(str).str.strip()


def norm_lower(s):
    return norm(s).str.lower()


def pct(n, d, decimals=1):
    if not d:
        return 0.0
    return round(100 * n / d, decimals)


def load_call_sheet(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[CALL_SHEET]
    data = ws.values
    header = next(data)
    df = pd.DataFrame(data, columns=header)
    return df, wb


def compute_funnel(df):
    A = norm(col(df, "A"))
    B = col(df, "B")
    E = norm(col(df, "E"))
    Qc = col(df, "Q")
    S = norm_lower(col(df, "S"))
    U = norm_lower(col(df, "U"))
    W = col(df, "W")

    total_leads = B.dropna().nunique()
    unique_attempts = Qc.notna().sum()
    connected = S.str.startswith("connected").sum()
    meaningful_connect = (S == "connected - spoke to customer").sum()
    interested = (
        (U == "interested - not yet (date known)").sum()
        + (U == "interested - not yet (date unknown)").sum()
    )
    booked = ((E == "1") & S.str.startswith("connected")).sum()
    booking_post_mc = ((E == "1") & (S == "connected - spoke to customer")).sum()

    f = dict(
        total_leads=int(total_leads),
        unique_attempts=int(unique_attempts),
        connected=int(connected),
        meaningful_connect=int(meaningful_connect),
        interested=int(interested),
        booked=int(booked),
        booking_post_mc=int(booking_post_mc),
    )
    f["attempts_pct"] = pct(f["unique_attempts"], f["total_leads"])
    f["connect_pct"] = pct(f["connected"], f["unique_attempts"])
    f["mc_pct"] = pct(f["meaningful_connect"], f["connected"])
    f["interest_pct"] = pct(f["interested"], f["meaningful_connect"])
    f["booked_pct"] = round(100 * f["booked"] / f["total_leads"], 1)

    f["w_attempts"] = round(100 * f["unique_attempts"] / f["total_leads"], 1)
    f["w_connected"] = round(100 * f["connected"] / f["total_leads"], 1)
    f["w_mc"] = round(100 * f["meaningful_connect"] / f["total_leads"], 1)
    f["w_interested"] = round(100 * f["interested"] / f["total_leads"], 1)
    f["w_booked"] = round(100 * f["booked"] / f["total_leads"], 1)

    f["arrow2"] = f["connect_pct"]
    f["arrow3"] = f["mc_pct"]
    f["arrow4"] = f["interest_pct"]
    f["arrow5"] = round(100 * f["booked"] / f["interested"], 1) if f["interested"] else 0.0

    # Follow-up summary: mirrors DAILY SUMMARY rows 4-7, which track the SECOND
    # calling round (columns AF:AS), not a simple filter of the first round.
    #   F4 = leads with a decision date on file (first-round W column)
    #   F5 = F4 - count of rows where AL (2nd-round "waiting for" note) is non-blank
    #   F6 = COUNTIFS(AI = "Connected - Spoke to customer")   [2nd-round connect status]
    #   F7 = COUNTIFS(AK = "Already Booked", AI = "Connected - Spoke to customer")
    fu_leads = int(B.loc[col(df, "W").notna()].dropna().nunique())
    AL_nonblank = col(df, "AL").notna().sum()
    fu_f5 = fu_leads - int(AL_nonblank)
    AI = norm_lower(col(df, "AI"))
    AK = norm_lower(col(df, "AK"))
    fu_mc = int((AI == "connected - spoke to customer").sum())
    fu_already_booked = int(((AK == "already booked") & (AI == "connected - spoke to customer")).sum())

    fu = dict(leads=fu_leads, mc=fu_mc)
    fu["attempts_pct"] = pct(fu_f5, fu_leads)
    fu["connect_pct"] = pct(fu_mc, fu_f5) if fu_f5 else 0.0
    fu["booked_pct"] = pct(fu_already_booked, fu_mc) if fu_mc else 0.0

    return f, fu


def compute_intent(df):
    A = norm(col(df, "A"))
    B = col(df, "B")
    E = norm(col(df, "E"))
    A_lower = A.str.lower()

    rows = []
    for label, color in INTENT_ORDER:
        mask = A_lower == label.lower()
        total = df.loc[mask].iloc[:, 1].dropna().nunique()
        booked = df.loc[mask & (E == "1")].iloc[:, 1].dropna().nunique()
        rows.append(dict(name=INTENT_DISPLAY_NAME.get(label, label), total=int(total), booked=int(booked), color=color))

    max_total = max(r["total"] for r in rows) or 1
    grand_total = sum(r["total"] for r in rows) or 1
    for r in rows:
        r["width"] = round(100 * r["total"] / max_total, 1)
        r["pct"] = round(100 * r["total"] / grand_total, 1)
    rows.sort(key=lambda r: -r["total"])
    return rows


def compute_decision_days(df):
    U = norm_lower(col(df, "U"))
    Z = pd.to_numeric(col(df, "Z"), errors="coerce")

    buckets = [
        ("Interested · date known", "interested - not yet (date known)", True),
    ]
    out = []
    for label, key, pill in buckets:
        mask = U == key
        n = int(mask.sum())
        vals = Z[mask].dropna()
        if len(vals) == 0:
            out.append(dict(label=label, n=n, median="N/A", avg="N/A", p90="N/A", p95="N/A", pill=pill))
        else:
            out.append(dict(
                label=label, n=n,
                median=round(vals.median(), 2) if vals.median() % 1 else int(vals.median()),
                avg=round(vals.mean(), 2),
                p90=round(vals.quantile(0.9), 2) if vals.quantile(0.9) % 1 else int(vals.quantile(0.9)),
                p95=round(vals.quantile(0.95), 2) if vals.quantile(0.95) % 1 else int(vals.quantile(0.95)),
                pill=pill,
            ))
    return out


def compute_outcome_median(df):
    A = norm(col(df, "A")).str.lower()
    Z = pd.to_numeric(col(df, "Z"), errors="coerce")

    def fmt(x):
        x = round(x, 1)
        return int(x) if x % 1 == 0 else x

    buckets = ["Wouldn't", "Not Yet"]
    out = []
    for label in buckets:
        mask = A == label.lower()
        vals = Z[mask].dropna()
        if len(vals) == 0:
            out.append(dict(label=label, median="N/A", avg="N/A", p90="N/A", p95="N/A"))
        else:
            out.append(dict(
                label=label,
                median=fmt(vals.median()),
                avg=fmt(vals.mean()),
                p90=fmt(vals.quantile(0.9)),
                p95=fmt(vals.quantile(0.95)),
            ))
    return out


def compute_reasons(wb, meaningful_connect):
    ws = wb[DAILY_SUMMARY]
    rows = []
    for r in range(70, 100):
        label = ws.cell(row=r, column=1).value  # A: friendly label
        count = ws.cell(row=r, column=3).value  # C: count
        bifurcation = ws.cell(row=r, column=7).value  # G: bifurcation
        if label is None or count in (None, "", 0):
            continue
        bifurcation = (bifurcation or "").strip() if isinstance(bifurcation, str) else bifurcation
        rows.append(dict(label=str(label).strip(), count=int(count), bifurcation=bifurcation))

    # top 5, excluding "Others" (own dict copies -- these labels also appear in
    # the breakdown groups below, which must not overwrite this "pct"/"width")
    ranked = sorted([dict(r) for r in rows if r["label"].lower() != "others"], key=lambda r: -r["count"])
    top5 = ranked[:5]
    for r in top5:
        r["pct"] = round(100 * r["count"] / meaningful_connect, 1)
    max_pct = top5[0]["pct"] if top5 else 1
    for r in top5:
        r["width"] = round(100 * r["pct"] / max_pct, 1)
    top5_sum = round(sum(r["pct"] for r in top5), 1)

    # breakdown groups
    grand_total = sum(r["count"] for r in rows if r["bifurcation"] in [b for b, _ in BREAKDOWN_BUCKETS])
    groups = []
    for bucket_name, color in BREAKDOWN_BUCKETS:
        members = [dict(r) for r in rows if r["bifurcation"] == bucket_name]
        members.sort(key=lambda r: -r["count"])
        total = sum(r["count"] for r in members)
        for r in members:
            r["pct"] = round(100 * r["count"] / grand_total, 1) if grand_total else 0.0
        groups.append(dict(
            name=BREAKDOWN_DISPLAY_NAME.get(bucket_name, bucket_name),
            color=color,
            total=total,
            share=round(100 * total / grand_total, 1) if grand_total else 0.0,
            reasons=members,
        ))

    return top5, top5_sum, groups, grand_total


def build(xlsx_path, output_path):
    xlsx_path = Path(xlsx_path)
    df, wb = load_call_sheet(xlsx_path)

    f, fu = compute_funnel(df)
    intent_categories = compute_intent(df)
    decision_buckets = compute_decision_days(df)
    outcome_median = compute_outcome_median(df)
    top5, top5_sum, breakdown_groups, grand_total = compute_reasons(wb, f["meaningful_connect"])
    voc_groups = classify_voc(df, col)

    ctx = dict(
        source_name=xlsx_path.stem,
        row_count=len(df),
        snapshot_date=datetime.now().strftime("%d %b %Y"),
        f=f,
        fu=fu,
        intent_categories=intent_categories,
        decision_buckets=decision_buckets,
        outcome_median=outcome_median,
        top5=top5,
        top5_sum=top5_sum,
        breakdown_groups=breakdown_groups,
        grand_total=grand_total,
        voc_groups=voc_groups,
    )

    template = Template(TEMPLATE_PATH.read_text(encoding="utf-8"))
    html = template.render(**ctx)
    Path(output_path).write_text(html, encoding="utf-8")
    print(f"Wrote {output_path}")
    print(f"  Total leads={f['total_leads']} Connected={f['connected']} MC={f['meaningful_connect']} Booked={f['booked']}")
    return ctx


if __name__ == "__main__":
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else None
    out_path = sys.argv[2] if len(sys.argv) > 2 else str(HERE / "dropoff_dashboard.html")
    if not xlsx_path:
        print("Usage: py build_dashboard.py <path-to-xlsx> [output-html-path]")
        sys.exit(1)
    build(xlsx_path, out_path)
